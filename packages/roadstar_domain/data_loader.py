"""Data loader for Hackathon_Data.xlsx."""
import os
import openpyxl
from datetime import datetime
from typing import List, Dict, Tuple, Optional
from .models import DriverModel, OrderModel, HOSStatus

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
EXCEL_DEFAULT_PATH = os.environ.get(
    "ROADSTAR_DATASET_PATH",
    os.path.join(BASE_DIR, "data", "Hackathon_Data.xlsx"),
)


def parse_dm_coord(coord_str: str) -> Optional[float]:
    """Parses coordinates like '0433201N' or '0795300W' into decimal degrees."""
    if not coord_str or str(coord_str).strip() in ["<null>", "", "None"]:
        return None
    s = str(coord_str).strip().upper()
    try:
        if s.endswith("N") or s.endswith("S"):
            hemi = s[-1]
            num = s[:-1]
            deg = float(num[:3])
            mins = float(num[3:5])
            secs = float(num[5:]) if len(num) > 5 else 0.0
            dec = deg + mins / 60.0 + secs / 3600.0
            return -dec if hemi == "S" else dec
        elif s.endswith("W") or s.endswith("E"):
            hemi = s[-1]
            num = s[:-1]
            deg = float(num[:3])
            mins = float(num[3:5])
            secs = float(num[5:]) if len(num) > 5 else 0.0
            dec = deg + mins / 60.0 + secs / 3600.0
            return -dec if hemi == "W" else dec
        return float(s)
    except Exception:
        return None


def load_dataset(excel_path: str = EXCEL_DEFAULT_PATH) -> Tuple[List[DriverModel], List[OrderModel]]:
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    
    # 1. Load Drivers
    driver_sheet = wb["Driver"]
    driver_rows = list(driver_sheet.iter_rows(values_only=True))
    d_headers = {h: idx for idx, h in enumerate(driver_rows[0]) if h}
    
    drivers: List[DriverModel] = []
    for row in driver_rows[1:]:
        if not row[d_headers.get("DRIVER_ID", 0)]:
            continue
        d_id = int(row[d_headers["DRIVER_ID"]])
        name = str(row[d_headers.get("FIRST_NAME", 5)] or f"Driver{d_id}")
        email = str(row[d_headers.get("EMAIL", 4)] or f"driver{d_id}@roadstar.ca")
        status = str(row[d_headers.get("STATUS", 35)] or "AVAIL")
        home_zone = str(row[d_headers.get("HOME_ZONE", 1)] or "RSTAR")
        lat = parse_dm_coord(str(row[d_headers.get("POSLAT", 31)]))
        lng = parse_dm_coord(str(row[d_headers.get("POSLONG", 32)]))
        loc_desc = str(row[d_headers.get("LAST_SAT_LOC", 30)] or "Southern Ontario")
        
        can_7 = float(row[d_headers.get("REMAINING_HOURS_CAN_7", 14)] or 55.0)
        can_14 = float(row[d_headers.get("REMAINING_HOURS_CAN_14", 16)] or 95.0)
        curr_duty = int(row[d_headers.get("CURRENT_DUTY", 19)] or 0)
        
        duty_desc = {0: "Off-Duty", 1: "Sleeper", 2: "Driving", 3: "On-Duty"}.get(curr_duty, "Off-Duty")
        
        hos = HOSStatus(
            driver_id=d_id,
            first_name=name,
            current_duty=curr_duty,
            duty_status_desc=duty_desc,
            remaining_hours_can_7=can_7,
            remaining_hours_can_14=can_14,
            shift_driving_hours=min(8.0, max(0.0, 13.0 - (can_7 % 13))),
            shift_on_duty_hours=min(9.5, max(0.0, 14.0 - (can_7 % 14))),
            shift_elapsed_hours=10.0,
            is_compliant=True
        )
        
        drivers.append(DriverModel(
            driver_id=d_id,
            first_name=name,
            email=email,
            home_zone=home_zone,
            status=status,
            lat=lat,
            lng=lng,
            location_desc=loc_desc,
            hos=hos
        ))

    # 2. Load Orders from Tlorder
    order_sheet = wb["Tlorder"]
    order_rows = list(order_sheet.iter_rows(values_only=True))
    o_headers = {h: idx for idx, h in enumerate(order_rows[0]) if h}
    rate_header = next((name for name in ("RATE_CAD", "RATE", "TOTAL_CHARGES", "FREIGHT_CHARGES")
                        if name in o_headers), None)
    
    orders: List[OrderModel] = []
    for row in order_rows[1:150]:  # Load active working batch
        bill_raw = row[o_headers.get("BILL_NUMBER", 1)]
        bill = str(bill_raw or "")
        if not bill or bill == "<null>":
            continue
        trip = str(row[o_headers.get("TRIP_NUMBER", 2)] or "")
        cust = str(row[o_headers.get("CALLNAME", 3)] or "")
        orig_city = str(row[o_headers.get("ORIGCITY", 4)] or "")
        orig_prov = str(row[o_headers.get("ORIGPROV", 5)] or "")
        orig_pc = str(row[o_headers.get("ORIGPC", 6)] or "")
        dest_city = str(row[o_headers.get("DESTCITY", 11)] or "")
        dest_prov = str(row[o_headers.get("DESTPROV", 12)] or "")
        dest_pc = str(row[o_headers.get("DESTPC", 13)] or "")
        dist_raw = row[o_headers.get("DISTANCE", 18)]
        weight_raw = row[o_headers.get("WEIGHT_LBS", 30)]
        rate_raw = row[o_headers[rate_header]] if rate_header else None
        # The source workbook currently has no freight-rate column. Never turn
        # a missing/zero source value into a plausible-looking commercial load.
        if (not all(v and str(v).strip() not in {"<null>", "None"}
                    for v in (cust, orig_city, orig_prov, dest_city, dest_prov))
                or dist_raw is None or weight_raw is None or rate_raw is None):
            continue
        dist = float(dist_raw)
        wgt = float(weight_raw)
        rate = float(rate_raw)
        if (dist <= 0 or wgt <= 0 or rate <= 0
                or (orig_city.strip().upper(), orig_prov.strip().upper())
                == (dest_city.strip().upper(), dest_prov.strip().upper())):
            continue
        load_type = str(row[o_headers.get("LOAD_TYPE", 28)] or "Dry Van")
        desc = str(row[o_headers.get("LOAD_DESCRIPTION", 29)] or "Manufactured Goods")
        pallets = int(row[o_headers.get("PALLETS", 31)] or 24)
        temp = str(row[o_headers.get("TEMPERATURE", 32)] or "Ambient")
        temp_ctrl = bool(row[o_headers.get("TEMP_CONTROLLED", 20)])
        pu_dt = row[o_headers.get("ACTUAL_PICKUP", 7)]
        del_dt = row[o_headers.get("ACTUAL_DELIVERY", 14)]
        assigned_flag = bool(row[o_headers.get("CURRENTLY_ASSIGNED", 25)])
        
        orders.append(OrderModel(
            bill_number=bill,
            trip_number=trip,
            customer_name=cust,
            origin_city=orig_city,
            origin_prov=orig_prov,
            origin_pc=orig_pc,
            dest_city=dest_city,
            dest_prov=dest_prov,
            dest_pc=dest_pc,
            distance_miles=dist,
            load_type=load_type,
            load_description=desc,
            weight_lbs=wgt,
            pallets=pallets,
            temperature=temp,
            temp_controlled=temp_ctrl,
            rate_cad=rate,
            pickup_time=pu_dt if isinstance(pu_dt, datetime) else None,
            delivery_time=del_dt if isinstance(del_dt, datetime) else None,
            currently_assigned=assigned_flag
        ))

    return drivers, orders
